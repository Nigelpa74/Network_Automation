import xml.etree.ElementTree as ET
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def limpiar_float(valor):
    """Convierte de forma segura un texto a float, ignorando 'NA', 'N/A', vacíos o textos no numéricos."""
    if valor is None:
        return None
    val_str = str(valor).strip()
    if val_str.upper() in ["NA", "N/A", "NULL", "-", ""]:
        return None
    try:
        return float(val_str)
    except ValueError:
        return None

def parsear_xml_sfp(ruta_o_contenido_xml):
    """
    Lee la estructura XML consolidada de NETCONF para extraer el inventario
    de puertos y potencias ópticas de todos los mininodos Huawei.
    """
    if isinstance(ruta_o_contenido_xml, str) and ruta_o_contenido_xml.strip().startswith("<"):
        root = ET.fromstring(ruta_o_contenido_xml)
    else:
        tree = ET.parse(ruta_o_contenido_xml)
        root = tree.getroot()

    namespaces = {
        'ns0': 'urn:ietf:params:xml:ns:netconf:base:1.0',
        'ns1': 'urn:huawei:yang:huawei-devm',
        'ns2': 'urn:huawei:yang:huawei-pic'
    }

    registros = []

    for router in root.findall('router'):
        nombre_router = router.get('nombre', '')
        ip_router = router.get('host', '')

        # Buscar todos los puertos dentro del árbol devm
        puertos = router.findall('.//ns1:port', namespaces)
        if not puertos:
            puertos = router.findall('.//port')

        for puerto in puertos:
            pos_node = puerto.find('ns1:position', namespaces)
            if pos_node is None:
                pos_node = puerto.find('position')
            posicion = pos_node.text if pos_node is not None else ''

            admin_node = puerto.find('ns1:admin-state', namespaces)
            if admin_node is None:
                admin_node = puerto.find('admin-state')
            estado_admin = admin_node.text if admin_node is not None else ''

            modulo_optico = puerto.find('ns2:optical-module', namespaces)
            if modulo_optico is None:
                for child in puerto:
                    if child.tag.endswith('optical-module'):
                        modulo_optico = child
                        break

            if modulo_optico is not None:
                sfp_presente = "Sí"
                
                wl_node = modulo_optico.find('ns2:wavelength', namespaces)
                if wl_node is None:
                    for child in modulo_optico:
                        if child.tag.endswith('wavelength'):
                            wl_node = child
                            break
                long_onda = limpiar_float(wl_node.text) if wl_node is not None else None

                dist_node = modulo_optico.find('ns2:transmission-distance', namespaces)
                if dist_node is None:
                    for child in modulo_optico:
                        if child.tag.endswith('transmission-distance'):
                            dist_node = child
                            break
                distancia = dist_node.text if dist_node is not None else ''

                # Búsqueda ultra-flexible de canal y potencias
                canal = None
                for elem in modulo_optico.iter():
                    if elem.tag.endswith('channel') and not elem.tag.endswith('channels'):
                        canal = elem
                        break

                if canal is not None:
                    tx_text, rx_text, bias_text = None, None, None
                    for sub in canal:
                        tag_clean = sub.tag.split('}')[-1] if '}' in sub.tag else sub.tag
                        if tag_clean == 'tx-power':
                            tx_text = sub.text
                        elif tag_clean == 'rx-power':
                            rx_text = sub.text
                        elif tag_clean == 'bias':
                            bias_text = sub.text
                    
                    pot_tx = limpiar_float(tx_text)
                    pot_rx = limpiar_float(rx_text)
                    bias = limpiar_float(bias_text)
                else:
                    pot_tx, pot_rx, bias = None, None, None
            else:
                sfp_presente = "No"
                long_onda = None
                distancia = "N/A"
                pot_tx, pot_rx, bias = None, None, None

            registros.append({
                "Nombre Router": nombre_router,
                "IP Host": ip_router,
                "Puerto": posicion,
                "Estado Admin": estado_admin,
                "SFP Presente": sfp_presente,
                "Long. Onda (nm)": long_onda,
                "Distancia / Especificación": distancia,
                "Potencia Tx (dBm)": pot_tx,
                "Potencia Rx (dBm)": pot_rx,
                "Corriente Bias (mA)": bias
            })

    return registros


def exportar_a_excel(registros, nombre_excel_salida="Inventario_Optico_Mininodos.xlsx"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reporte DDM Mininodos"

    ws.views.sheetView[0].showGridLines = True

    ws.merge_cells("A1:J1")
    title_cell = ws["A1"]
    title_cell.value = "INVENTARIO DE TRANSCEPTORES Y NIVELES ÓPTICOS (DDM) - MININODOS MINEDU"
    title_cell.font = Font(name="Segoe UI", size=13, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:J2")
    sub_cell = ws["A2"]
    sub_cell.value = "Datos consolidados vía NETCONF / YANG (huawei-devm & huawei-pic)"
    sub_cell.font = Font(name="Segoe UI", size=10, italic=True, color="595959")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 18

    cabeceras = [
        "Nombre Router", "IP Host", "Puerto", "Estado Admin", "SFP Presente",
        "Long. Onda (nm)", "Distancia / Especificación", "Potencia Tx (dBm)", "Potencia Rx (dBm)", "Corriente Bias (mA)"
    ]

    ws.append([])
    ws.append(cabeceras)

    header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    borde_fino = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    ws.row_dimensions[4].height = 26

    for col_num, header_text in enumerate(cabeceras, 1):
        cell = ws.cell(row=4, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = borde_fino

    for idx, reg in enumerate(registros, start=5):
        ws.append([
            reg["Nombre Router"],
            reg["IP Host"],
            reg["Puerto"],
            reg["Estado Admin"],
            reg["SFP Presente"],
            reg["Long. Onda (nm)"],
            reg["Distancia / Especificación"],
            reg["Potencia Tx (dBm)"],
            reg["Potencia Rx (dBm)"],
            reg["Corriente Bias (mA)"]
        ])
        ws.row_dimensions[idx].height = 20

        sombreado = PatternFill(start_color="F9FBFD" if idx % 2 == 0 else "FFFFFF", fill_type="solid")

        for col_idx in range(1, 11):
            cell = ws.cell(row=idx, column=col_idx)
            cell.fill = sombreado
            cell.border = borde_fino
            cell.font = Font(name="Segoe UI", size=9.5)

            if col_idx in [1, 2, 3, 7]:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif col_idx in [4, 5]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="right", vertical="center")

            if col_idx == 6 and cell.value is not None:
                cell.number_format = "#,##0"
            elif col_idx in [8, 9, 10] and cell.value is not None:
                cell.number_format = "0.00"

    ws.auto_filter.ref = f"A4:J{len(registros) + 4}"

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["F"].width = 18
    ws.column_dimensions["G"].width = 28
    ws.column_dimensions["H"].width = 18
    ws.column_dimensions["I"].width = 18
    ws.column_dimensions["J"].width = 20

    wb.save(nombre_excel_salida)
    print(f"Excel generado con éxito: {nombre_excel_salida}")

if __name__ == "__main__":
    ARCHIVO_XML_ENTRADA = "resultado_inventario_sfp_devm.xml"
    ARCHIVO_EXCEL_SALIDA = "Inventario_Optico_Mininodos.xlsx"

    print("Procesando XML de mininodos...")
    try:
        datos_extraidos = parsear_xml_sfp(ARCHIVO_XML_ENTRADA)
        exportar_a_excel(datos_extraidos, ARCHIVO_EXCEL_SALIDA)
    except FileNotFoundError:
        print(f"Archivo no encontrado: {ARCHIVO_XML_ENTRADA}")