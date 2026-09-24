import os
import xml.etree.ElementTree as ET
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def parse_netconf_xml(xml_source, hostname="rAcH1_Loreto5 (Nauta)"):
    """Parsea el XML de NETCONF de Huawei y extrae los datos de puertos y SFP."""
    # Si el parámetro es una ruta de archivo existente, parsea el archivo; si no, asume string XML
    if os.path.exists(xml_source):
        tree = ET.parse(xml_source)
        root = tree.getroot()
    else:
        root = ET.fromstring(xml_source)

    # Namespaces presentes en la respuesta de Huawei
    namespaces = {
        "nc": "urn:ietf:params:xml:ns:netconf:base:1.0",
        "devm": "urn:huawei:yang:huawei-devm",
        "pic": "urn:huawei:yang:huawei-pic",
    }

    ports_data = []

    # Buscar todos los elementos <port> dentro del namespace devm
    for port in root.findall(".//devm:port", namespaces):
        pos_elem = port.find("devm:position", namespaces)
        admin_elem = port.find("devm:admin-state", namespaces)
        opt_elem = port.find("pic:optical-module", namespaces)

        position = pos_elem.text if pos_elem is not None else "N/A"
        admin_state = admin_elem.text if admin_elem is not None else "N/A"

        # Clasificación simple de tipo de interfaz
        if "100GE" in position:
            iftype = "100G Ethernet"
        elif "10GE" in position:
            iftype = "10G Ethernet"
        elif "GE" in position:
            iftype = "1G Ethernet"
        else:
            iftype = "Management"

        if opt_elem is not None:
            has_sfp = "Sí"
            pn_elem = opt_elem.find("pic:vendor-pn", namespaces)
            wl_elem = opt_elem.find("pic:wavelength", namespaces)
            dist_elem = opt_elem.find("pic:transmission-distance", namespaces)
            mfg_elem = opt_elem.find("pic:manufacture-date", namespaces)
            sn_elem = opt_elem.find("pic:serial-number", namespaces)

            vendor_pn = pn_elem.text if pn_elem is not None else "N/A"
            wavelength = (
                f"{wl_elem.text} nm"
                if wl_elem is not None and wl_elem.text
                else "N/A"
            )
            raw_dist = (
                dist_elem.text if dist_elem is not None and dist_elem.text else ""
            )

            # Normalización básica de la distancia (ej. 20000(9um...) -> 20 km)
            if "20000" in raw_dist:
                distance = "20 km"
            elif "40000" in raw_dist:
                distance = "40 km"
            elif "10000" in raw_dist:
                distance = "10 km"
            else:
                distance = raw_dist if raw_dist else "N/A"

            mfg_date = mfg_elem.text if mfg_elem is not None else "N/A"
            serial_number = sn_elem.text if sn_elem is not None else "N/A"
        else:
            has_sfp = "No"
            vendor_pn = "N/A"
            wavelength = "N/A"
            distance = "N/A"
            mfg_date = "N/A"
            serial_number = "N/A"

        ports_data.append(
            {
                "Nodo": hostname,
                "Puerto": position,
                "Tipo Interfaz": iftype,
                "Admin State": admin_state,
                "SFP Instalado": has_sfp,
                "Part Number (PN)": vendor_pn,
                "Serial Number (SN)": serial_number,
                "Longitud Onda": wavelength,
                "Distancia Nomina": distance,
                "Fecha Fabricación": mfg_date,
            }
        )

    return ports_data


def generate_excel_report(data, output_filename="Inventario_Optico.xlsx"):
    """Genera un archivo Excel formateado con openpyxl a partir de los datos parseados."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventario_Transceptores"

    # Encabezados
    headers = list(data[0].keys())
    ws.append(headers)

    # Estilos
    header_fill = PatternFill(
        start_color="1F497D", end_color="1F497D", fill_type="solid"
    )  # Azul oscuro corporativo
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    # Rellenos sutiles para filas
    no_sfp_fill = PatternFill(
        start_color="F2F2F2", end_color="F2F2F2", fill_type="solid"
    )  # Gris para vacíos

    # Formatear Encabezados
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    # Agregar filas de datos con formato
    for row_num, row_data in enumerate(data, 2):
        for col_num, key in enumerate(headers, 1):
            val = row_data[key]
            cell = ws.cell(row=row_num, column=col_num, value=val)
            cell.font = data_font
            cell.border = thin_border

            # Alineación según el tipo de campo
            if key in [
                "Puerto",
                "Admin State",
                "SFP Instalado",
                "Longitud Onda",
                "Distancia Nomina",
                "Fecha Fabricación",
            ]:
                cell.alignment = center_align
            else:
                cell.alignment = left_align

            # Sombreado para puertos sin SFP
            if row_data["SFP Instalado"] == "No":
                cell.fill = no_sfp_fill

    # Activar Autofiltro en la tabla
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(data) + 1}"

    # Auto-ajustar ancho de columnas
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(output_filename)
    print(f"Reporte generado exitosamente: {output_filename}")


# --- EJECUCIÓN PRINCIPAL ---
if __name__ == "__main__":
    archivo_xml = "rAcH1_Loreto5 (Nauta)_devm.xml"

    # Parsear y generar
    datos_extraidos = parse_netconf_xml(archivo_xml)
    generate_excel_report(datos_extraidos, "Inventario_Optico_Nauta.xlsx")