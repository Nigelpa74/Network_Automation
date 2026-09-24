import re
import xml.etree.ElementTree as ET
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def parsear_xml_consolidado(xml_path):
    tree = ET.parse(xml_path)
    root = tree.getroot()

    todos_los_puertos = []

    # Iterar por cada nodo <router> del XML
    for router in root.findall("router"):
        nombre_nodo = router.attrib.get("nombre", "Desconocido")
        host_ip = router.attrib.get("host", "N/A")

        # Buscar todos los elementos <port> ignorando el namespace exacto con {*}
        for port in router.findall(".//{*}port"):
            pos_elem = port.find("{*}position")
            admin_elem = port.find("{*}admin-state")
            opt_elem = port.find("{*}optical-module")

            position = (
                pos_elem.text if pos_elem is not None and pos_elem.text else "N/A"
            )
            admin_state = (
                admin_elem.text
                if admin_elem is not None and admin_elem.text
                else "N/A"
            )

            # Clasificación del tipo de interfaz por nombre de puerto
            if "100GE" in position:
                iftype = "100G Ethernet"
            elif "10GE" in position:
                iftype = "10G Ethernet"
            elif "GE" in position:
                iftype = "1G Ethernet"
            else:
                iftype = "Management"

            if opt_elem is not None:
                has_sfp = "Sí"
                pn_elem = opt_elem.find("{*}vendor-pn")
                wl_elem = opt_elem.find("{*}wavelength")
                dist_elem = opt_elem.find("{*}transmission-distance")
                mfg_elem = opt_elem.find("{*}manufacture-date")
                sn_elem = opt_elem.find("{*}serial-number")
                desc_elem = opt_elem.find("{*}description")  # 👈 1. Extraer etiqueta description

                vendor_pn = (
                    pn_elem.text
                    if pn_elem is not None and pn_elem.text
                    else "N/A"
                )
                wavelength = (
                    f"{wl_elem.text} nm"
                    if wl_elem is not None and wl_elem.text
                    else "N/A"
                )

                raw_dist = (
                    dist_elem.text
                    if dist_elem is not None and dist_elem.text
                    else ""
                )

                # Limpieza de la distancia (ej. "40000(9um/125um SMF)" -> "40 km")
                match_dist = re.search(r"(\d+)", raw_dist)
                if match_dist:
                    meters = int(match_dist.group(1))
                    distance = (
                        f"{meters // 1000} km"
                        if meters >= 1000
                        else f"{meters} m"
                    )
                else:
                    distance = raw_dist if raw_dist else "N/A"

                mfg_date = (
                    mfg_elem.text
                    if mfg_elem is not None and mfg_elem.text
                    else "N/A"
                )
                serial_number = (
                    sn_elem.text
                    if sn_elem is not None and sn_elem.text
                    else "N/A"
                )
                sfp_desc = (
                    desc_elem.text
                    if desc_elem is not None and desc_elem.text
                    else "N/A"
                )  # 👈 Valor de velocidad/descripción
            else:
                has_sfp = "No"
                vendor_pn = "N/A"
                wavelength = "N/A"
                distance = "N/A"
                mfg_date = "N/A"
                serial_number = "N/A"
                sfp_desc = "N/A"

            # 👈 2. Se agrega la clave 'Velocidad/Descripción SFP' al reporte
            todos_los_puertos.append(
                {
                    "Nodo": nombre_nodo,
                    "IP Host": host_ip,
                    "Puerto": position,
                    "Tipo Interfaz": iftype,
                    "Admin State": admin_state,
                    "SFP Instalado": has_sfp,
                    "Velocidad/Descripción SFP": sfp_desc,
                    "Part Number (PN)": vendor_pn,
                    "Serial Number (SN)": serial_number,
                    "Longitud Onda": wavelength,
                    "Distancia Nóminal": distance,
                    "Fecha Fabricación": mfg_date,
                }
            )

    return todos_los_puertos


def generar_excel(datos_puertos, output_filename="Reporte_SFP_Consolidado.xlsx"):
    if not datos_puertos:
        print("❌ No se encontraron datos para exportar.")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "Inventario_Puertos"

    headers = list(datos_puertos[0].keys())
    ws.append(headers)

    # Definición de estilos
    header_fill = PatternFill(
        start_color="1F497D", end_color="1F497D", fill_type="solid"
    )
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    no_sfp_fill = PatternFill(
        start_color="F9F9F9", end_color="F9F9F9", fill_type="solid"
    )

    # Aplicar estilos a encabezados
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    # Agregar filas y aplicar formatos
    for row_num, row_data in enumerate(datos_puertos, 2):
        for col_num, key in enumerate(headers, 1):
            val = row_data[key]
            cell = ws.cell(row=row_num, column=col_num, value=val)
            cell.font = data_font
            cell.border = thin_border

            # 👈 3. Se añade 'Velocidad/Descripción SFP' a la lista de alineación centrada
            if key in [
                "IP Host",
                "Puerto",
                "Admin State",
                "SFP Instalado",
                "Velocidad/Descripción SFP",
                "Longitud Onda",
                "Distancia Nóminal",
                "Fecha Fabricación",
            ]:
                cell.alignment = center_align
            else:
                cell.alignment = left_align

            # Sombrear ligeramente las filas donde no hay SFP instalado
            if row_data["SFP Instalado"] == "No":
                cell.fill = no_sfp_fill

    # Agregar auto-filtros y ajustar ancho de columnas
    ws.auto_filter.ref = (
        f"A1:{get_column_letter(len(headers))}{len(datos_puertos) + 1}"
    )
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    wb.save(output_filename)
    print(f"📊 [EXCEL GENERADO] Reporte guardado en: {output_filename}")


if __name__ == "__main__":
    archivo_xml = "resultado_inventario_sfp_devm.xml"
    puertos = parsear_xml_consolidado(archivo_xml)
    generar_excel(puertos, "Inventario_Optico_Huawei.xlsx")